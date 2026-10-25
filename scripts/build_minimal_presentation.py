"""Build an editable, locally reproducible 16-slide presentation from saved results."""

import csv
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from fraudetection.experiments import MODEL_LABELS, MODEL_ORDER, PARAMETER_GRIDS
from fraudetection.formatting import es_number

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "documentacion/presentacion/canva/presentacion.pptx"
PAPER = "FFFFFF"
INK = "171717"
TEAL = "2563EB"
MUTED = "64748B"
PALE = "EFF6FF"
RULE = "E2E8F0"
WHITE = "FFFFFF"
CORAL = "B91C1C"
FONT = "Arial"


def read_csv(path):
    with (ROOT / path).open(encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def color(value):
    return RGBColor.from_string(value)


def textbox(
    slide,
    x,
    y,
    w,
    h,
    text,
    *,
    size=18,
    col=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
    margin=0,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = str(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color(col)
    return shape


def rect(slide, x, y, w, h, fill, *, line=None, radius=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.fill.background() if line is None else None
    if line is not None:
        shape.line.color.rgb = color(line)
        shape.line.width = Pt(0.7)
    style = shape._element.find(qn("p:style"))
    if style is not None:
        effect = style.find(qn("a:effectRef"))
        if effect is not None:
            effect.set("idx", "0")
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def line(slide, x1, y1, x2, y2, col=RULE, width=1):
    s = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    s.line.color.rgb = color(col)
    s.line.width = Pt(width)
    return s


def add_slide(prs, index, section, title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(PAPER)
    textbox(slide, 0.65, 0.30, 5.8, 0.22, section.upper(), size=9, col=TEAL, bold=True)
    textbox(slide, 0.65, 0.63, 11.9, 0.62, title, size=27, bold=True)
    if subtitle:
        textbox(slide, 0.67, 1.27, 11.8, 0.40, subtitle, size=12, col=MUTED)
    line(slide, 0.65, 7.12, 12.68, 7.12, RULE, 0.8)
    textbox(
        slide,
        0.67,
        7.17,
        9.8,
        0.20,
        "INFO1185 · Detección de fraude · datos y resultados del proyecto",
        size=8,
        col=MUTED,
    )
    textbox(
        slide, 11.75, 7.15, 0.85, 0.24, f"{index:02d} / 16", size=8, col=MUTED, align=PP_ALIGN.RIGHT
    )
    return slide


def metric(slide, x, y, w, value, label, detail=""):
    textbox(slide, x, y, w, 0.65, value, size=24 if w < 3 else 30, col=TEAL, bold=True)
    textbox(slide, x, y + 0.65, w, 0.33, label, size=12, bold=True)
    if detail:
        textbox(slide, x, y + 1.0, w, 0.47, detail, size=10, col=MUTED, valign=MSO_ANCHOR.TOP)


def bar_chart(
    slide,
    rows,
    value_key,
    *,
    x,
    y,
    w,
    h,
    value_label="",
    std_key=None,
    value_scale=100,
    show_labels=True,
    color_key=None,
):
    left = x + 1.55
    chart_w = w - 3.15
    row_h = h / len(rows)
    for i, row in enumerate(rows):
        cy = y + i * row_h
        label = row.get("short", row.get("model_label", ""))
        textbox(
            slide,
            x,
            cy + 0.03,
            1.45,
            row_h - 0.04,
            label,
            size=10,
            col=INK,
            valign=MSO_ANCHOR.MIDDLE,
        )
        value = float(row[value_key]) * 100
        bw = chart_w * min(value / value_scale, 1)
        rect(
            slide,
            left,
            cy + row_h * 0.27,
            bw,
            row_h * 0.42,
            (TEAL if row.get(color_key, "") == "True" else INK) if color_key else TEAL,
        )
        val = f"{es_number(value, 1)}%"
        if std_key:
            val += f" ± {es_number(float(row[std_key]) * 100, 1)} pp"
        textbox(
            slide,
            left + chart_w + 0.08,
            cy + 0.01,
            1.65,
            row_h - 0.02,
            val,
            size=10,
            col=MUTED,
        )
    textbox(
        slide,
        left,
        y + h + 0.02,
        chart_w,
        0.18,
        f"{value_label} · escala 0–{value_scale}%",
        size=8,
        col=MUTED,
        align=PP_ALIGN.RIGHT,
    )


def table(
    slide, headers, rows, *, x, y, widths, row_h=0.48, font=10, emphasis_col=None, header_font=9
):
    total = sum(widths)
    rect(slide, x, y, total, row_h, INK)
    cx = x
    for i, h in enumerate(headers):
        textbox(
            slide,
            cx + 0.08,
            y + 0.02,
            widths[i] - 0.12,
            row_h - 0.04,
            h,
            size=header_font,
            col=WHITE,
            bold=True,
        )
        cx += widths[i]
    for ri, row in enumerate(rows):
        yy = y + row_h * (ri + 1)
        rect(slide, x, yy, total, row_h, WHITE if ri % 2 == 0 else PALE)
        cx = x
        for ci, cell in enumerate(row):
            is_emph = emphasis_col == ci
            textbox(
                slide,
                cx + 0.08,
                yy + 0.02,
                widths[ci] - 0.12,
                row_h - 0.04,
                cell,
                size=font,
                col=TEAL if is_emph else INK,
                bold=is_emph,
            )
            cx += widths[ci]


def pct(value, std=None, digits=1, std_digits=None):
    result = f"{es_number(float(value) * 100, digits)}%"
    if std is not None:
        precision = digits if std_digits is None else std_digits
        result += f" ± {es_number(float(std) * 100, precision)} pp"
    return result


def params(raw):
    values = json.loads(raw)
    return ", ".join(
        f"{key.removeprefix('model__')}={grid_value(value)}" for key, value in values.items()
    )


def grid_value(value):
    if value is None:
        return "libre"
    if isinstance(value, float):
        return es_number(value, 1 if value % 1 else 0)
    return str(value)


def run():
    cv = read_csv("resultados/cv_resumen.csv")
    test = read_csv("resultados/test_resumen.csv")
    preds = read_csv("resultados/predicciones_configuracion_cv.csv")
    metadata = json.loads((ROOT / "resultados/metadata.json").read_text(encoding="utf-8"))
    expected_conditions = {(model, flag) for model in MODEL_ORDER for flag in ("False", "True")}
    for name, rows in (("CV", cv), ("prueba", test)):
        actual = {(row["model"], row["compensated"]) for row in rows}
        if len(rows) != 8 or actual != expected_conditions:
            raise ValueError(f"{name}: se requieren los cuatro modelos en ambas condiciones")
    if sum(row["selected_by_cv"] == "True" for row in cv) != 1:
        raise ValueError("CV debe identificar exactamente una configuración ganadora")
    if len(preds) != metadata["test_rows"]:
        raise ValueError("Las predicciones no coinciden con el tamaño del test")
    by_exp = {r["experiment"]: r for r in test}
    selected = next(r for r in cv if r["selected_by_cv"] == "True")
    selected_test = by_exp[selected["experiment"]]
    cm = {(int(p["class_true"]), int(p["class_predicted"])): 0 for p in preds}
    for p in preds:
        key = (int(p["class_true"]), int(p["class_predicted"]))
        cm[key] = cm.get(key, 0) + 1

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Detección de fraude con clasificadores supervisados"
    prs.core_properties.author = "Fernando Alonso Valdés Gómez; Juan Francisco Muñoz Veloso"
    prs.core_properties.subject = "Respaldo local reproducible basado en resultados guardados"

    # 1. Portada
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color(PAPER)
    rect(s, 0, 0, 0.22, 7.5, TEAL)
    textbox(s, 0.9, 0.65, 4, 0.3, "INFO1185 · PROYECTO", size=11, col=TEAL, bold=True)
    textbox(
        s, 0.9, 1.55, 8.7, 1.6, "Detectar fraude\ncon aprendizaje supervisado", size=34, bold=True
    )
    textbox(
        s,
        0.94,
        3.48,
        8,
        0.5,
        "Comparación de ocho condiciones con sensibilidad como criterio principal",
        size=17,
        col=MUTED,
    )
    line(s, 0.95, 4.5, 6.7, 4.5, TEAL, 2)
    textbox(s, 0.95, 4.78, 6.8, 0.42, "Fernando Alonso Valdés Gómez", size=16)
    textbox(s, 0.95, 5.22, 6.8, 0.42, "Juan Francisco Muñoz Veloso", size=16)
    textbox(s, 9.15, 1.65, 3.0, 0.6, "99,827%", size=26, col=TEAL, bold=True)
    textbox(s, 9.18, 2.28, 3.2, 0.34, "exactitud al predecir todo legítimo", size=11, col=MUTED)
    textbox(s, 9.15, 3.1, 3.0, 0.6, "0%", size=26, col=CORAL, bold=True)
    textbox(s, 9.18, 3.73, 3.2, 0.34, "sensibilidad para fraude", size=11, col=MUTED)
    line(s, 0.95, 7.12, 12.68, 7.12, RULE, 0.8)
    textbox(s, 0.95, 7.17, 10, 0.2, "Proyecto de INFO1185 · 2026", size=8, col=MUTED)
    textbox(s, 11.75, 7.15, 0.85, 0.24, "01 / 16", size=8, col=MUTED, align=PP_ALIGN.RIGHT)

    # 2. Problema
    s = add_slide(
        prs,
        2,
        "Contexto",
        "La exactitud puede ocultar el fraude",
        "La clase positiva es la transacción fraudulenta.",
    )
    metric(
        s, 0.9, 2.05, 3.2, "99,827%", "Exactitud base", "284.315 de 284.807 filas son legítimas."
    )
    metric(
        s,
        4.9,
        2.05,
        3.2,
        "0%",
        "Sensibilidad base",
        "Predecir siempre legítimo omite los 492 fraudes.",
    )
    rect(s, 9.05, 2.1, 3.2, 2.15, PALE)
    textbox(s, 9.35, 2.34, 2.6, 0.45, "Objetivo", size=12, col=TEAL, bold=True)
    textbox(
        s,
        9.35,
        2.87,
        2.55,
        1.0,
        "Encontrar más fraude y vigilar falsas alarmas.",
        size=18,
        bold=True,
    )
    textbox(
        s,
        0.9,
        5.15,
        10.8,
        0.7,
        (
            "La evaluación prioriza sensibilidad. La precisión muestra cuántas alertas "
            "corresponden a fraude."
        ),
        size=17,
    )

    # 3. Datos y limpieza
    s = add_slide(
        prs,
        3,
        "Datos",
        "Una clase rara y duplicados exactos",
        "Credit Card Fraud Detection · CSV inspeccionado antes del modelado.",
    )
    metric(s, 0.9, 2.0, 2.7, "284.807", "filas originales")
    metric(s, 3.9, 2.0, 2.7, "492", "fraudes originales", "0,173% de las filas.")
    metric(s, 6.9, 2.0, 2.7, "1.081", "duplicados exactos")
    metric(s, 9.9, 2.0, 2.7, "283.726", "filas tras limpieza")
    line(s, 0.9, 4.2, 12.3, 4.2)
    textbox(s, 0.95, 4.48, 5.3, 0.4, "30 predictores: Time, Amount y V1–V28", size=15, bold=True)
    textbox(
        s,
        0.95,
        4.98,
        5.3,
        0.78,
        "Class excluida · 0 faltantes · 473 fraudes tras limpiar. "
        "Se quitan duplicados antes de separar train y test.",
        size=12,
        col=MUTED,
    )
    textbox(s, 6.6, 4.48, 5.4, 0.4, "No aplicamos nuevo PCA.", size=17, bold=True)
    textbox(
        s,
        6.6,
        4.98,
        5.3,
        0.78,
        "V1–V28 ya son componentes del dataset.",
        size=13,
        col=MUTED,
    )

    # 4. Particiones
    s = add_slide(
        prs,
        4,
        "Diseño experimental",
        "Particiones fijas y prueba reservada",
        "La selección se hace en train. El test se usa al final.",
    )
    rect(s, 0.95, 2.05, 5.35, 2.3, PALE)
    textbox(
        s,
        1.25,
        2.3,
        4.7,
        0.4,
        f"TRAIN · 70% estratificado · semilla {metadata['cv_seed']}",
        size=15,
        col=TEAL,
        bold=True,
    )
    textbox(
        s,
        1.25,
        2.9,
        4.7,
        0.66,
        f"{es_number(metadata['train_rows'])} filas",
        size=27,
        bold=True,
    )
    textbox(
        s,
        1.25,
        3.62,
        4.7,
        0.35,
        (
            f"Submuestra: {es_number(metadata['search_sample_rows'])} · "
            f"{es_number(metadata['search_sample_fraud'])} fraudes"
        ),
        size=12,
        col=MUTED,
    )
    rect(s, 7.0, 2.05, 5.35, 2.3, INK)
    textbox(
        s,
        7.3,
        2.3,
        4.7,
        0.4,
        "TEST · 30% estratificado · reservado",
        size=15,
        col="B8D8D1",
        bold=True,
    )
    textbox(
        s,
        7.3,
        2.9,
        4.7,
        0.66,
        f"{es_number(metadata['test_rows'])} filas",
        size=27,
        col=WHITE,
        bold=True,
    )
    textbox(s, 7.3, 3.62, 4.7, 0.35, "142 fraudes · 84.976 legítimas", size=12, col="D5DEDA")
    textbox(
        s,
        1.1,
        5.05,
        10.8,
        0.6,
        (
            f"Validación cruzada estratificada de {metadata['cv_folds']} pliegues. "
            "Los ocho experimentos comparten los mismos pliegues."
        ),
        size=16,
    )

    textbox(
        s,
        1.1,
        5.9,
        10.8,
        0.65,
        f"Búsqueda con {es_number(metadata['search_sample_rows'])} filas; "
        f"reajuste final sobre las {es_number(metadata['train_rows'])} filas de entrenamiento, "
        "con el submuestreo correspondiente a cada condición.",
        size=14,
        col=MUTED,
    )

    # 5. Experimentos
    s = add_slide(
        prs,
        5,
        "Método",
        "Ocho condiciones, una comparación controlada",
        "Dos condiciones: sin compensación y submuestreo de 1 fraude por cada 10 legítimas.",
    )
    transformations = {
        "knn": "StandardScaler · ANOVA F top 20",
        "tree": "ANOVA F top 20",
        "svm": "StandardScaler · ANOVA F top 20",
        "forest": "ANOVA F top 20",
    }
    experiments = [
        [MODEL_LABELS[m], "Sin compensación · submuestreo 1:10", transformations[m]]
        for m in MODEL_ORDER
    ]
    table(
        s,
        ["Clasificador", "Condiciones", "Transformación"],
        experiments,
        x=0.8,
        y=1.9,
        widths=[2.75, 3.65, 4.4],
        row_h=0.67,
        font=12,
        header_font=10,
    )
    textbox(
        s,
        0.9,
        5.45,
        5.55,
        1.25,
        "Cada transformación y remuestreo se ajusta solo con el train de cada pliegue: "
        "sin fuga de información.",
        size=12,
        col=MUTED,
        valign=MSO_ANCHOR.TOP,
    )
    textbox(
        s,
        6.65,
        5.45,
        5.6,
        1.35,
        "Seleccionamos las 20 variables con mayor puntuación ANOVA F, un criterio marginal. "
        "El submuestreo conserva fraudes, pero descarta legítimas. "
        "Escalar importa para distancias y regularización.",
        size=12,
        col=MUTED,
        valign=MSO_ANCHOR.TOP,
    )

    # 6. Random Forest
    s = add_slide(
        prs,
        6,
        "Método",
        "Random Forest propio: árboles diversos por bootstrap",
        (
            "Número de árboles, profundidad, max_features y semilla configurables; "
            "el ganador usa 21 árboles y profundidad 6."
        ),
    )
    steps = [
        (
            "01",
            "Bootstrap",
            "Muestra con reemplazo, del tamaño del train de cada condición; "
            "ya submuestreado cuando corresponde.",
        ),
        (
            "02",
            "Árbol de decisión diverso",
            (
                "Semillas distintas derivadas de la semilla general; "
                "max_features='sqrt' en cada división."
            ),
        ),
        (
            "03",
            "Legítima o fraude",
            "predict suma los votos de los árboles; "
            "el empate favorece fraude. fit y predict están separados.",
        ),
    ]
    for i, (num, heading, body) in enumerate(steps):
        xx = 0.9 + i * 4.05
        textbox(s, xx, 2.15, 0.7, 0.45, num, size=15, col=TEAL, bold=True)
        line(s, xx, 2.74, xx + 3.3, 2.74, TEAL, 1.6)
        textbox(s, xx, 2.95, 3.5, 0.45, heading, size=17, bold=True)
        textbox(s, xx, 3.55, 3.55, 0.95, body, size=13, col=MUTED, valign=MSO_ANCHOR.TOP)
    textbox(
        s,
        0.95,
        5.45,
        10.9,
        0.6,
        (
            "Bagging reduce la variabilidad; max_features por división "
            "reduce la correlación entre árboles."
        ),
        size=17,
    )

    # 7. Búsqueda
    s = add_slide(
        prs,
        7,
        "Selección",
        "La búsqueda prioriza sensibilidad y desempata por precisión",
        "Grillas acotadas por costo; C se muestra con coma decimal.",
    )
    grid_rows = []
    for model in MODEL_ORDER:
        grid = PARAMETER_GRIDS[model]
        grid_rows.append(
            [
                MODEL_LABELS[model],
                "; ".join(
                    f"{k.removeprefix('model__')}: "
                    f"{'/'.join(grid_value(value) for value in values)}"
                    for k, values in grid.items()
                ),
            ]
        )
    table(
        s,
        ["Modelo", "Valores explorados"],
        grid_rows,
        x=1.0,
        y=1.95,
        widths=[2.8, 7.4],
        row_h=0.64,
        font=12,
    )
    rect(s, 1.0, 5.4, 10.8, 0.9, PALE)
    textbox(
        s,
        1.3,
        5.55,
        10.0,
        0.6,
        (
            f"Ganador CV: {selected['model_label']} con {selected['compensation']} · "
            f"{params(selected['best_parameters'])} · sensibilidad "
            f"{pct(selected['cv_sensitivity_mean'], selected['cv_sensitivity_std'])}"
        ),
        size=16,
        col=TEAL,
        bold=True,
    )

    textbox(
        s,
        1.0,
        6.38,
        10.8,
        0.56,
        (
            "Vecinos: suavidad; profundidad y hojas: complejidad; C: regularización; "
            "árboles: diversidad y costo. Pesos por distancia fijos en k-NN."
        ),
        size=11,
        col=MUTED,
    )

    # 8. CV chart
    s = add_slide(
        prs,
        8,
        "Resultados · validación",
        "La sensibilidad más alta aparece con Random Forest compensado",
        f"Media ± desviación estándar entre los {metadata['cv_folds']} pliegues.",
    )
    rows = [
        {
            **r,
            "short": (
                f"{r['model_label']} · {'1:10' if r['compensated'] == 'True' else 'sin comp.'}"
            ),
        }
        for r in cv
    ]
    bar_chart(
        s,
        rows,
        "cv_sensitivity_mean",
        x=0.9,
        y=1.9,
        w=11.3,
        h=4.45,
        value_label="Sensibilidad CV",
        std_key="cv_sensitivity_std",
        value_scale=100,
    )
    textbox(
        s,
        1.0,
        6.45,
        10.8,
        0.32,
        "Mayor media observada; no demuestra superioridad estadística. La precisión desempata.",
        size=10,
        col=MUTED,
    )

    # 9. Test chart
    s = add_slide(
        prs,
        9,
        "Resultados · prueba",
        "La sensibilidad final se estima con cinco reajustes",
        "Mismo conjunto de prueba en cada repetición; se muestran sensibilidad y FNR.",
    )
    textbox(s, 3.1, 1.84, 4.6, 0.28, "Sensibilidad · escala 0–100%", size=10, col=TEAL, bold=True)
    textbox(s, 8.0, 1.84, 3.8, 0.28, "FNR · escala 0–50%", size=10, col=CORAL, bold=True)
    row_h = 0.51
    for i, row in enumerate(test):
        yy = 2.18 + i * row_h
        short = f"{row['model_label']} · {'1:10' if row['compensated'] == 'True' else 'sin comp.'}"
        textbox(s, 0.82, yy, 2.15, 0.38, short, size=9, col=INK)
        sens = float(row["sensitivity_mean"]) * 100
        fnr = float(row["fnr_mean"]) * 100
        rect(s, 3.1, yy + 0.09, 3.35 * sens / 100, 0.22, TEAL)
        textbox(
            s,
            6.52,
            yy,
            1.25,
            0.38,
            f"{es_number(sens, 1)} ± {es_number(float(row['sensitivity_std']) * 100, 1)}",
            size=8,
            col=MUTED,
        )
        rect(s, 8.0, yy + 0.09, 2.4 * min(fnr / 50, 1), 0.22, CORAL)
        textbox(
            s,
            10.47,
            yy,
            1.48,
            0.38,
            f"{es_number(fnr, 1)} ± {es_number(float(row['fnr_std']) * 100, 1)}",
            size=8,
            col=MUTED,
        )
    textbox(
        s,
        1.0,
        6.43,
        10.8,
        0.32,
        "Medias en %; desviaciones estándar en puntos porcentuales (pp).",
        size=10,
        col=MUTED,
    )

    # 10. Tree vs forest
    s = add_slide(
        prs,
        10,
        "Resultados · compensación",
        "Árbol de decisión y Random Forest muestran el costo de recuperar fraude",
        "Promedios en test; falsas alarmas por repetición.",
    )
    compare = []
    for model in ("tree", "forest"):
        for comp in ("False", "True"):
            r = next(x for x in test if x["model"] == model and x["compensated"] == comp)
            compare.append(
                [
                    f"{r['model_label']} · {'1:10' if comp == 'True' else 'sin comp.'}",
                    pct(r["sensitivity_mean"]),
                    pct(r["precision_mean"]),
                    es_number(float(r["fp_mean"]), 1),
                ]
            )
    table(
        s,
        ["Condición", "Sensibilidad", "Precisión", "FP promedio"],
        compare,
        x=0.95,
        y=2.15,
        widths=[4.0, 2.0, 2.0, 2.0],
        row_h=0.7,
        font=15,
        header_font=10,
    )
    textbox(
        s,
        1.0,
        5.65,
        10.7,
        0.7,
        "Con 1:10, Random Forest tiene menor sensibilidad que el árbol de decisión, pero menos FP.",
        size=15,
        col=TEAL,
        bold=True,
    )

    textbox(
        s,
        1.0,
        6.2,
        10.8,
        0.35,
        "El ganador se mantuvo según CV; no se cambió tras mirar el test.",
        size=12,
        col=MUTED,
    )

    # 11. Confusion matrix from saved prediction rows
    s = add_slide(
        prs,
        11,
        "Resultados · semilla 42",
        "Matriz de confusión del ganador en un ajuste",
        "Predicciones guardadas para la configuración elegida por CV.",
    )
    matrix = [[cm[(0, 0)], cm[(0, 1)]], [cm[(1, 0)], cm[(1, 1)]]]
    x0, y0 = 3.2, 2.05
    textbox(
        s, 3.3, 1.8, 3.4, 0.25, "Predicción legítimo", size=11, col=MUTED, align=PP_ALIGN.CENTER
    )
    textbox(s, 6.8, 1.8, 3.4, 0.25, "Predicción fraude", size=11, col=MUTED, align=PP_ALIGN.CENTER)
    for ri, label in enumerate(("Real legítimo", "Real fraude")):
        textbox(s, 0.9, y0 + ri * 1.35 + 0.4, 2.0, 0.42, label, size=12, col=MUTED, bold=True)
        for ci in range(2):
            xx, yy = x0 + ci * 3.55, y0 + ri * 1.35
            fill = PALE if ri == ci else ("F3E4DE" if ci == 1 else PALE)
            rect(s, xx, yy, 3.1, 1.1, fill)
            textbox(
                s,
                xx,
                yy + 0.1,
                3.1,
                0.6,
                es_number(matrix[ri][ci]),
                size=27,
                col=TEAL if ri == ci else CORAL,
                bold=True,
                align=PP_ALIGN.CENTER,
            )
            textbox(
                s,
                xx,
                yy + 0.7,
                3.1,
                0.25,
                ("TN" if ri == 0 and ci == 0 else "FP" if ri == 0 else "FN" if ci == 0 else "TP"),
                size=10,
                col=MUTED,
                align=PP_ALIGN.CENTER,
            )
    textbox(
        s,
        1.0,
        5.45,
        10.7,
        0.5,
        f"Captura {matrix[1][1]} de 142 fraudes; genera {matrix[0][1]} falsas alarmas.",
        size=17,
        bold=True,
    )
    textbox(
        s,
        1.0,
        6.03,
        10.7,
        0.34,
        "Una sola semilla explica la matriz; el resumen final promedia cinco semillas.",
        size=11,
        col=MUTED,
    )

    # 12. Conclusion
    s = add_slide(
        prs,
        12,
        "Cierre",
        "Ganador seleccionado por validación: Random Forest con submuestreo",
        "La selección se fijó antes de consultar el test.",
    )
    metric(
        s,
        0.75,
        2.05,
        2.65,
        pct(selected["cv_sensitivity_mean"], selected["cv_sensitivity_std"]),
        "Sensibilidad CV",
        f"Media ± DE · {metadata['cv_folds']} pliegues.",
    )
    metric(
        s,
        3.75,
        2.05,
        2.65,
        pct(selected_test["sensitivity_mean"], selected_test["sensitivity_std"]),
        "Sensibilidad test",
        "Media ± DE · cinco reajustes.",
    )
    metric(
        s,
        6.75,
        2.05,
        2.65,
        pct(selected_test["precision_mean"], selected_test["precision_std"]),
        "Precisión test",
        "Alertas que son fraude.",
    )
    metric(
        s,
        9.75,
        2.05,
        2.65,
        pct(selected_test["fnr_mean"], selected_test["fnr_std"]),
        "FNR test",
        "Fracción de fraude omitido.",
    )
    textbox(s, 1.0, 4.55, 10.9, 0.42, "Límites", size=16, bold=True)
    textbox(
        s,
        1.0,
        5.0,
        10.9,
        1.0,
        "Test con 142 fraudes; las cinco repeticiones comparten los casos. ANOVA F es "
        "marginal, la grilla es acotada y no se hizo validación temporal.",
        size=14,
        col=MUTED,
    )

    # 13. CV table
    s = add_slide(
        prs,
        13,
        "Respaldo · validación",
        "Parámetros y métricas CV de los ocho experimentos",
        "Medias en %; desviaciones estándar en pp. `weights=distance` se mantiene fijo en k-NN.",
    )
    rows = []
    for r in cv:
        rows.append(
            [
                r["model_label"],
                "1:10" if r["compensated"] == "True" else "Ninguna",
                params(r["best_parameters"]),
                pct(r["cv_sensitivity_mean"], r["cv_sensitivity_std"]),
                pct(r["cv_precision_mean"]),
            ]
        )
    table(
        s,
        ["Modelo", "Muestreo", "Mejores parámetros", "Sens. CV", "Prec. CV"],
        rows,
        x=0.55,
        y=1.85,
        widths=[2.15, 1.35, 4.45, 2.1, 1.65],
        row_h=0.55,
        font=10,
        header_font=8,
    )

    # 14/15 test detail
    for compensate, idx, label in (
        ("False", 14, "Sin compensación"),
        ("True", 15, "Con submuestreo 1:10"),
    ):
        s = add_slide(
            prs,
            idx,
            "Respaldo · prueba",
            f"Seis métricas · {label}",
            "Media en % ± desviación estándar en pp · cinco reajustes sobre el mismo test.",
        )
        subset = [r for r in test if r["compensated"] == compensate]
        rows = []
        for r in subset:
            rows.append(
                [
                    r["model_label"],
                    pct(r["sensitivity_mean"], r["sensitivity_std"], digits=1),
                    pct(r["specificity_mean"], r["specificity_std"], digits=3, std_digits=3),
                    pct(r["precision_mean"], r["precision_std"], digits=1),
                    pct(r["f1_mean"], r["f1_std"], digits=1),
                    pct(r["accuracy_mean"], r["accuracy_std"], digits=2, std_digits=2),
                    pct(r["fnr_mean"], r["fnr_std"], digits=1),
                ]
            )
        table(
            s,
            ["Modelo", "Sens.", "Espec.", "Prec.", "F1", "Exact.", "FNR"],
            rows,
            x=0.53,
            y=2.05,
            widths=[2.25, 1.75, 1.75, 1.75, 1.65, 1.8, 1.55],
            row_h=0.72,
            font=10,
            header_font=9,
        )
        textbox(
            s,
            0.65,
            6.0,
            11.6,
            0.38,
            (
                "Espec. con 3 decimales y exactitud con 2 para no ocultar FP; "
                "DE en puntos porcentuales."
            ),
            size=10,
            col=MUTED,
        )

    # 16. Baselines and next steps
    s = add_slide(
        prs,
        16,
        "Respaldo · interpretación",
        "Dos reglas simples muestran el costo del desbalance",
        "Cálculos sobre el test principal, con 84.976 legítimas y 142 fraudes.",
    )
    rect(s, 0.9, 2.05, 5.35, 2.05, PALE)
    textbox(s, 1.2, 2.28, 4.7, 0.35, "Siempre legítimo", size=18, bold=True)
    textbox(
        s, 1.2, 2.82, 4.7, 0.38, "Exactitud 99,833% · sensibilidad 0%", size=14, col=TEAL, bold=True
    )
    textbox(s, 1.2, 3.35, 4.7, 0.38, "142 fraudes omitidos", size=13, col=MUTED)
    rect(s, 7.0, 2.05, 5.35, 2.05, "F3E4DE")
    textbox(s, 7.3, 2.28, 4.7, 0.35, "Siempre fraude", size=18, bold=True)
    textbox(
        s,
        7.3,
        2.82,
        4.7,
        0.38,
        "Sensibilidad 100% · especificidad 0%",
        size=14,
        col=CORAL,
        bold=True,
    )
    textbox(s, 7.3, 3.35, 4.7, 0.38, "84.976 falsas alarmas", size=13, col=MUTED)
    line(s, 0.95, 4.65, 12.25, 4.65)
    textbox(s, 1.0, 4.95, 5.0, 0.35, "Por qué se busca sensibilidad", size=15, bold=True)
    textbox(
        s,
        1.0,
        5.4,
        5.0,
        0.78,
        (
            "El costo de omitir fraude puede ser alto; la precisión y los FP miden "
            "la carga de revisión."
        ),
        size=12,
        col=MUTED,
    )
    textbox(s, 6.7, 4.95, 5.1, 0.35, "Siguientes mejoras", size=15, bold=True)
    textbox(
        s,
        6.7,
        5.4,
        5.1,
        0.86,
        (
            "Probar umbrales, calibración, validación temporal y nuevos datos "
            "solo con un test nuevo y separado."
        ),
        size=12,
        col=MUTED,
    )

    assert len(prs.slides) == 16, "La presentación requiere 12 diapositivas y cuatro respaldos"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Created {OUT.relative_to(ROOT)} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    run()
