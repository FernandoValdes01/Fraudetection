import json

import matplotlib.pyplot as plt
import pandas as pd
import pytest

from fraudetection import presentation


@pytest.mark.parametrize("existing_metadata", [False, True])
def test_build_report_uses_cv_metadata_and_preserves_protocol(
    tmp_path, monkeypatch, existing_metadata
):
    results_dir = tmp_path / "results"
    results_dir.mkdir()
    dataset = tmp_path / "dataset.csv"
    dataset.write_text("synthetic dataset")

    pd.DataFrame(
        {
            "experiment": ["exp-1", "exp-2"],
            "selected_by_cv": [True, False],
            "search_rows": [10_000, 10_000],
            "cv_folds": [4, 4],
        }
    ).to_csv(results_dir / "cv_resumen.csv", index=False)
    pd.DataFrame(
        {
            "experiment": ["exp-1"],
            "compensation": ["ninguna"],
            "sensitivity_mean": [0.8],
            "fnr_mean": [0.2],
        }
    ).to_csv(results_dir / "test_resumen.csv", index=False)
    pd.DataFrame({"class_true": [0, 1], "class_predicted": [0, 1]}).to_csv(
        results_dir / "predicciones_configuracion_cv.csv", index=False
    )
    original_metadata = {
        "cv_seed": 42,
        "search_sample_fraud": 17,
        "otrosprotocolo": {"test_repetitions": 3},
    }
    if existing_metadata:
        (results_dir / "metadata.json").write_text(json.dumps(original_metadata))

    frame = pd.DataFrame({"Class": [0, 0, 1, 1]})
    monkeypatch.setattr(presentation, "load_dataset", lambda _: frame)
    monkeypatch.setattr(
        presentation,
        "split_dataset",
        lambda clean: (None, pd.DataFrame(index=range(2)), None, None),
    )
    monkeypatch.setattr(presentation, "save_plots", lambda *args: None)
    monkeypatch.setattr(
        presentation,
        "create_deck",
        lambda cv, test, *args: (cv.iloc[0], test.iloc[0]),
    )

    metadata = presentation.build_report(dataset, results_dir, tmp_path / "deck.pptx")
    persisted = json.loads((results_dir / "metadata.json").read_text())

    assert metadata["search_sample_rows"] == 10_000
    assert metadata["cv_folds"] == 4
    assert persisted == metadata
    if existing_metadata:
        assert {key: persisted[key] for key in original_metadata} == original_metadata


@pytest.mark.parametrize("search_rows", [[10_000, 20_000], [10_000, None], [0, 0]])
def test_build_report_rejects_inconsistent_search_protocol(tmp_path, monkeypatch, search_rows):
    pd.DataFrame({"search_rows": search_rows, "cv_folds": [5, 5]}).to_csv(
        tmp_path / "cv_resumen.csv", index=False
    )
    frame = pd.DataFrame({"Class": [0, 1]})
    monkeypatch.setattr(presentation, "load_dataset", lambda _: frame)
    monkeypatch.setattr(presentation, "split_dataset", lambda _: (None, frame, None, None))
    with pytest.raises(ValueError, match="search_rows"):
        presentation.build_report(tmp_path / "dataset.csv", tmp_path, tmp_path / "deck.pptx")
    assert not (tmp_path / "metadata.json").exists()
    assert not (tmp_path / "deck.pptx").exists()


def test_spanish_number_format_matches_the_beamer_presentation():
    assert presentation.es_number(284807) == "284.807"
    assert presentation.es_number(84923) == "84.923"
    assert presentation.es_number(492) == "492"
    assert presentation.es_number(0.17275, 3) == "0,173"
    assert presentation.fmt(0.8042253) == "80,4 %"
    assert presentation.fmt(0.99996, 3) == "99,996 %"


def test_metrics_with_std_uses_spanish_separators():
    row = {"sensitivity_mean": 0.8042253, "sensitivity_std": 0.0031493}
    assert presentation.metric_with_std(row, "sensitivity") == "80,4 % ± 0,3 pp"


def test_describe_search_space_follows_the_real_grid():
    text = presentation.describe_search_space()
    assert "n_neighbors=3, 5, 7" in text
    assert "max_depth=5, 10, libre" in text
    assert "C=0,1, 1, 10" in text
    assert "n_estimators=11, 21" in text


def _write_plot(path):
    figure, axes = plt.subplots()
    axes.plot([0, 1], [0, 1])
    figure.savefig(path)
    plt.close(figure)


def test_deck_shows_the_derived_figures(tmp_path):
    results_dir = tmp_path / "resultados"
    plots = results_dir / "graficos"
    plots.mkdir(parents=True)
    for name in (
        "distribucion_clases",
        "sensibilidad_validacion",
        "sensibilidad_fnr_prueba",
        "precision_especificidad_prueba",
        "matriz_confusion_seleccion_cv",
    ):
        _write_plot(plots / f"{name}.png")

    cv = pd.DataFrame(
        {
            "experiment": ["forest_submuestreo", "tree_submuestreo"],
            "model": ["forest", "tree"],
            "model_label": ["Random Forest propio", "Árbol de decisión"],
            "compensation": ["submuestreo 1:10", "submuestreo 1:10"],
            "compensated": [True, True],
            "search_rows": [50_000, 50_000],
            "cv_folds": [5, 5],
            "best_parameters": ['{"max_depth": 6}', '{"max_depth": 5}'],
            "cv_sensitivity_mean": [0.8786764, 0.8544117],
            "cv_sensitivity_std": [0.0552205, 0.0641688],
            "cv_precision_mean": [0.4499389, 0.1857732],
            "selected_by_cv": [True, False],
        }
    )
    cv.to_csv(results_dir / "cv_resumen.csv", index=False)

    test = pd.DataFrame(
        {
            "experiment": ["forest_submuestreo", "forest_sin_compensacion", "tree_submuestreo"],
            "model": ["forest", "forest", "tree"],
            "model_label": ["Random Forest propio", "Random Forest propio", "Árbol de decisión"],
            "compensation": ["submuestreo 1:10", "ninguna", "submuestreo 1:10"],
            "compensated": [True, False, True],
            "test_repeats": [5, 5, 5],
            "sensitivity_mean": [0.8042253, 0.7619718, 0.8309859],
            "sensitivity_std": [0.0031493, 0.0115716, 0.0272745],
            "specificity_mean": [0.9993456, 0.9999270, 0.9969214],
            "specificity_std": [0.0001403, 0.0000153, 0.0013786],
            "precision_mean": [0.6751120, 0.9459713, 0.3357661],
            "precision_std": [0.0457185, 0.0101788, 0.1013236],
            "f1_mean": [0.7333611, 0.8439675, 0.4699230],
            "f1_std": [0.0276777, 0.0055296, 0.0961042],
            "accuracy_mean": [0.9990201, 0.9995300, 0.9966446],
            "accuracy_std": [0.0001408, 0.0000143, 0.0013375],
            "fnr_mean": [0.1957746, 0.2380281, 0.1690140],
            "fnr_std": [0.0031493, 0.0115716, 0.0272745],
        }
    )
    test.to_csv(results_dir / "test_resumen.csv", index=False)
    predictions = pd.DataFrame(
        {
            "class_true": [0] * 10_000 + [1] * 1_000,
            "class_predicted": [0] * 9_500 + [1] * 500 + [1] * 1_000,
        }
    )
    predictions.to_csv(results_dir / "predicciones_configuracion_cv.csv", index=False)

    figures = presentation.DeckFigures(
        raw_records=284_807,
        raw_frauds=492,
        raw_fraud_share=492 / 284_807,
        duplicates=1_081,
        missing_values=0,
        clean_records=283_726,
        clean_frauds=473,
        train_rows=198_608,
        test_rows=85_118,
        test_frauds=142,
        test_negatives=84_976,
    )
    deck_path = tmp_path / "deck.pptx"
    presentation.create_deck(cv, test, results_dir, deck_path, figures)

    from pptx import Presentation

    text = "\n".join(
        shape.text_frame.text
        for slide in Presentation(deck_path).slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert "284.807" in text
    assert "0,173 %" in text
    assert "84.976" in text
    assert "9.500" in text
    assert "80,4 %" in text
    assert "87,9 % ± 5,5 %" in text
    assert "2,7 puntos porcentuales menos" in text
    assert "Fernando Alonso Valdés Gómez" in text
    assert "9,500" not in text
