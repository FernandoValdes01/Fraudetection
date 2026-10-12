"""Gráficos científicos y documentos derivados exclusivamente de resultados guardados."""

import base64
import html
import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from fraudetection.config import AUTHORS
from fraudetection.experiments import MODEL_LABELS
from fraudetection.metrics import METRIC_NAMES

METRIC_LABELS = {
    "sensitivity": "Sensibilidad",
    "specificity": "Especificidad",
    "precision": "Precisión",
    "f1": "F1",
    "accuracy": "Exactitud",
    "fnr": "Tasa FN",
}
BLUE = "#24566c"
ORANGE = "#ce6b38"


def load_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def label(row):
    condition = "con" if row["condition"] == "con_compensacion" else "sin"
    return f"{MODEL_LABELS[row['model']]} · {condition} compensación"


def percent(mean, std):
    return f"{100 * mean:.2f} ± {100 * std:.2f}"


def metric_table(frame, metrics):
    rows = []
    for _, row in frame.iterrows():
        rows.append(
            [
                label(row),
                *[percent(row[f"{metric}_mean"], row[f"{metric}_std"]) for metric in metrics],
            ]
        )
    return pd.DataFrame(rows, columns=["Experimento", *[METRIC_LABELS[m] for m in metrics]])


def make_figures(output, exploration, cv, summary, repetitions):
    figures = output / "graficos"
    figures.mkdir(exist_ok=True)
    plt.rcParams.update(
        {
            "font.size": 11,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "figure.dpi": 120,
            "savefig.dpi": 160,
        }
    )
    fig, axes = plt.subplots(1, 2, figsize=(11, 4.5), layout="constrained")
    counts = exploration["sample_class_counts"]
    axes[0].bar(["Legítimas", "Fraudes"], [counts["0"], counts["1"]], color=[BLUE, ORANGE])
    axes[0].set_yscale("log")
    axes[0].set_ylabel("Transacciones, escala logarítmica")
    axes[0].set_title("Muestra estratificada")
    for i, value in enumerate(counts.values()):
        axes[0].text(i, value * 1.15, f"{value:,}", ha="center")
    axes[1].bar(
        ["Train legítimas", "Train fraude", "Test legítimas", "Test fraude"],
        [*exploration["train_class_counts"].values(), *exploration["test_class_counts"].values()],
        color=[BLUE, ORANGE, BLUE, ORANGE],
    )
    axes[1].set_yscale("log")
    axes[1].tick_params(axis="x", labelrotation=25)
    axes[1].set_title("Partición estratificada 70/30")
    fig.savefig(figures / "distribucion.png")
    plt.close(fig)

    for name, frame, title in [
        ("validacion", cv, "Validación: cinco pliegues"),
        ("prueba", summary, "Prueba: tres reajustes, mismo test"),
    ]:
        fig, axes = plt.subplots(1, 2, figsize=(12, 5), layout="constrained")
        positions = np.arange(len(frame))
        labels = [label(row).replace(" compensación", "") for _, row in frame.iterrows()]
        colors = [
            ORANGE if row["condition"] == "con_compensacion" else BLUE
            for _, row in frame.iterrows()
        ]
        for axis, metric in zip(axes, ("sensitivity", "precision"), strict=True):
            axis.barh(
                positions,
                100 * frame[f"{metric}_mean"],
                xerr=100 * frame[f"{metric}_std"],
                color=colors,
                capsize=3,
            )
            axis.set_yticks(positions, labels)
            axis.invert_yaxis()
            axis.set_xlim(0, 108)
            axis.set_xlabel(f"{METRIC_LABELS[metric]} (%)")
        fig.suptitle(title)
        fig.savefig(figures / f"{name}.png")
        plt.close(fig)

    fig, axes = plt.subplots(1, 2, figsize=(12, 5), layout="constrained")
    for axis, column, title in zip(
        axes, ("tp", "fp"), ("Fraudes detectados", "Falsas alarmas"), strict=True
    ):
        positions = np.arange(4)
        for offset, condition, color in [
            (-0.18, "sin_compensacion", BLUE),
            (0.18, "con_compensacion", ORANGE),
        ]:
            subset = (
                summary[summary.condition == condition].set_index("model").loc[list(MODEL_LABELS)]
            )
            axis.bar(
                positions + offset,
                subset[f"{column}_mean"],
                width=0.36,
                yerr=subset[f"{column}_std"],
                capsize=3,
                color=color,
                label=condition.replace("_", " "),
            )
        axis.set_xticks(positions, list(MODEL_LABELS.values()), rotation=15)
        axis.set_title(title)
        axis.set_ylabel("Media de transacciones")
        axis.legend(fontsize=8)
    fig.savefig(figures / "costos.png")
    plt.close(fig)

    first_seed = int(repetitions.seed.iloc[0])
    first = repetitions[repetitions.seed == first_seed]
    for _, row in first.iterrows():
        matrix = np.array([[row.tn, row.fp], [row.fn, row.tp]], dtype=int)
        fig, axis = plt.subplots(figsize=(4, 3.5), layout="constrained")
        axis.imshow(np.log1p(matrix), cmap="Blues")
        for (i, j), value in np.ndenumerate(matrix):
            axis.text(
                j,
                i,
                str(value),
                ha="center",
                va="center",
                fontsize=16,
                color="white" if np.log1p(value) > np.log1p(matrix.max()) / 2 else "black",
            )
        axis.set_xticks([0, 1], ["Legítima", "Fraude"])
        axis.set_yticks([0, 1], ["Legítima", "Fraude"])
        axis.set_xlabel("Predicción")
        axis.set_ylabel("Clase real")
        axis.set_title(f"{label(row)}\nSemilla {first_seed}", fontsize=10)
        fig.savefig(figures / f"matriz_{row['model']}_{row['condition']}.png")
        plt.close(fig)


def analysis_points(summary, selected, cv, exploration):
    points = []
    for model in MODEL_LABELS:
        rows = summary[summary.model == model].set_index("condition")
        plain, balanced = rows.loc["sin_compensacion"], rows.loc["con_compensacion"]
        delta_tp = balanced.tp_mean - plain.tp_mean
        delta_fp = balanced.fp_mean - plain.fp_mean
        points.append(
            f"{MODEL_LABELS[model]}: compensar cambia los fraudes detectados en {delta_tp:+.2f} "
            f"y las falsas alarmas en {delta_fp:+.2f}, medias de los tres ajustes."
        )
    winner = selected["winner"]
    chosen_cv = cv[(cv.model == winner["model"]) & (cv.condition == winner["condition"])].iloc[0]
    chosen_test = summary[
        (summary.model == winner["model"]) & (summary.condition == winner["condition"])
    ].iloc[0]
    winner_text = (
        f"Elegido por validación: {label(winner)}. Sensibilidad CV "
        f"{percent(chosen_cv.sensitivity_mean, chosen_cv.sensitivity_std)} %; "
        "sensibilidad test "
        f"{percent(chosen_test.sensitivity_mean, chosen_test.sensitivity_std)} %; "
        f"precisión test {percent(chosen_test.precision_mean, chosen_test.precision_std)} %."
    )
    comparisons = []
    for condition in ("sin_compensacion", "con_compensacion"):
        tree = summary[(summary.model == "tree") & (summary.condition == condition)].iloc[0]
        forest = summary[(summary.model == "forest") & (summary.condition == condition)].iloc[0]
        condition_label = (
            "Con compensación" if condition == "con_compensacion" else "Sin compensación"
        )
        comparisons.append(
            f"{condition_label}: bosque frente a árbol cambia la sensibilidad en "
            f"{100 * (forest.sensitivity_mean - tree.sensitivity_mean):+.2f} puntos porcentuales "
            f"y las falsas alarmas en {forest.fp_mean - tree.fp_mean:+.2f}."
        )
    limitations = [
        f"El test contiene {exploration['test_class_counts']['1']} fraudes. "
        "Pocos positivos hacen sensibles las métricas a cada error individual.",
        "La desviación en test mide variación del ajuste sobre un test fijo; "
        "no estima incertidumbre sobre otras muestras ni es un intervalo de confianza.",
        "La muestra reduce el costo computacional, pero descarta fraudes y limita la comparación. "
        "El submuestreo también elimina legítimas potencialmente informativas.",
        "La partición aleatoria exigida no valida cambios temporales ni otros bancos. "
        "ANOVA es un criterio marginal y puede descartar interacciones útiles.",
        "El bosque reduce correlación al combinar bootstrap y variables aleatorias; "
        "sus beneficios deben juzgarse por los resultados, sin suponer una mejora automática.",
    ]
    return points, winner_text, comparisons, limitations


def image_tag(path, alt):
    encoded = base64.b64encode(path.read_bytes()).decode()
    return f'<img alt="{html.escape(alt)}" src="data:image/png;base64,{encoded}">'


def make_html(output, exploration, cv, summary, selected, points, winner_text, comparisons, limits):
    protocol = load_json(output / "protocolo.json")
    figures = output / "graficos"
    parts = [
        """<!doctype html><html lang="es"><meta charset="utf-8">
    <meta name="viewport" content="width=device-width,initial-scale=1">
    <title>Fraudetection · resultados reproducibles</title><style>
    body{margin:0;background:#f4f3ee;color:#173440;font:16px/1.6 system-ui,sans-serif}
    main{max-width:1160px;margin:auto;padding:48px 24px}h1{font-size:48px;line-height:1.1}
    h2{margin-top:40px}img{width:100%;height:auto;background:white;border-radius:8px}
    .kpis{display:flex;gap:24px;flex-wrap:wrap}.kpis p{background:white;padding:18px;flex:1}
    .table-wrap{overflow:auto}table{border-collapse:collapse;width:100%;font-size:13px}
    th,td{padding:10px;border-bottom:1px solid #c9d0d1;text-align:left;white-space:nowrap}
    th{background:#173440;color:white}li{margin:8px 0}.note{background:#fff;padding:20px}
    @media print{body{background:white}main{padding:0}h1{font-size:32px}}
    </style><main><p>INFO1185 · Inteligencia Artificial</p>
    <h1>Detectar fraude.<br>Medir lo que se pierde.</h1>"""
    ]
    parts.append(f'<p class="note">{html.escape(winner_text)}</p>')
    parts.append(
        '<div class="kpis">'
        + "".join(
            f"<p>{html.escape(name)}<br><strong>{value:,}</strong></p>"
            for name, value in (
                ("Filas de la muestra", exploration["sample_rows"]),
                ("Fraudes en entrenamiento", exploration["train_class_counts"]["1"]),
                ("Fraudes en prueba", exploration["test_class_counts"]["1"]),
            )
        )
        + "</div>"
    )
    parts.append("<h2>Datos y metodología</h2>")
    parts.append(
        f"<p>Fuente: Credit Card Fraud Detection, Kaggle, copia pública usada por "
        f"<a href='https://www.tensorflow.org/tutorials/structured_data/imbalanced_data'>"
        f"TensorFlow</a>. {exploration['raw_rows']:,} filas originales; "
        f"{exploration['duplicate_rows_removed']:,} duplicados exactos eliminados. "
        "Sin faltantes. Muestra estratificada sin reemplazo y partición 70/30, semilla 42. "
        "Se conservan 30 predictores originales y se excluye Class.</p>"
    )
    parts.append(image_tag(figures / "distribucion.png", "Distribución de las clases"))
    parts.append(
        "<p>Cinco pliegues estratificados compartidos. Escalado para k-NN y SVM; "
        "selección ANOVA de 20 características dentro de cada pliegue, antes del "
        "submuestreo aleatorio opcional con razón fraude/legítima 0,1. "
        "Todos los candidatos se eligen por sensibilidad y se desempata por precisión. "
        "El test conserva su prevalencia y no participa en decisiones.</p>"
    )
    parts.append(
        "<h2>Random Forest propio</h2><p>31 árboles independientes; cada bootstrap "
        "muestrea n filas con reemplazo. Semillas distintas por árbol derivadas de "
        "la general. Cada división considera sqrt variables y la etiqueta final se "
        "obtiene por mayoría de votos. Solo se utiliza DecisionTreeClassifier como base. "
        "Profundidad 8 o libre; mínimo de dos observaciones por hoja.</p>"
    )
    parts.append("<h2>Validación y configuraciones elegidas</h2>")
    params = cv[["model", "condition", "params"]].copy()
    params["model"] = params.model.map(MODEL_LABELS)
    parts.append('<div class="table-wrap">' + params.to_html(index=False, border=0) + "</div>")
    parts.append("<p>Valores en porcentaje: media ± desviación entre cinco pliegues, ddof=0.</p>")
    parts.append(
        '<div class="table-wrap">'
        + metric_table(cv, METRIC_NAMES).to_html(index=False, border=0)
        + "</div>"
    )
    parts.append(image_tag(figures / "validacion.png", "Sensibilidad y precisión en validación"))
    parts.append(
        "<h2>Prueba reservada</h2><p>Tres reajustes con semillas 42, 43 y 44, mismas "
        "configuraciones y mismo test. Porcentajes, media ± desviación muestral, ddof=1. "
        "En procedimientos deterministas la desviación puede ser cero.</p>"
    )
    parts.append(
        '<div class="table-wrap">'
        + metric_table(summary, METRIC_NAMES).to_html(index=False, border=0)
        + "</div>"
    )
    parts.append(image_tag(figures / "prueba.png", "Sensibilidad y precisión en prueba"))
    parts.append("<h2>Fraudes y falsas alarmas</h2>")
    parts.append(image_tag(figures / "costos.png", "Efecto de compensar el desbalance"))
    parts.append("<ul>" + "".join(f"<li>{html.escape(p)}</li>" for p in points) + "</ul>")
    parts.append(
        "<h2>Árbol frente a bosque</h2><ul>"
        + "".join(f"<li>{html.escape(p)}</li>" for p in comparisons)
        + "</ul>"
    )
    parts.append(
        "<h2>Matrices de confusión</h2><p>Conteos de la semilla 42; las tres "
        "repeticiones se encuentran en prueba_repeticiones.csv.</p>"
    )
    for path in sorted(figures.glob("matriz_*.png")):
        parts.append(
            '<div style="display:inline-block;width:min(100%,480px)">'
            + image_tag(path, path.stem)
            + "</div>"
        )
    baselines = load_json(output / "referencias_triviales.json")
    legitimate = baselines["always_legitimate"]
    fraud = baselines["always_fraud"]
    parts.append(
        f"<h2>Por qué la exactitud engaña</h2><p>Predecir siempre legítimo alcanza "
        f"{100 * legitimate['accuracy']:.3f} % de exactitud, con sensibilidad 0 % y "
        f"{legitimate['fn']} fraudes omitidos. Predecir siempre fraude alcanza "
        f"sensibilidad 100 %, pero produce {fraud['fp']:,} falsas alarmas.</p>"
    )
    parts.append(
        "<h2>Límites y conclusiones</h2><ul>"
        + "".join(f"<li>{html.escape(p)}</li>" for p in limits)
        + "</ul>"
    )
    parts.append(
        f"<p>Duración de experimentos: {protocol['elapsed_seconds'] / 60:.2f} minutos. "
        "Las semillas, versiones, grillas y particiones quedan registradas. "
        "No se modificó el protocolo tras observar prueba. "
        "Este experimento académico no valida un sistema para uso bancario.</p>"
    )
    parts.append("</main></html>")
    (output / "analisis.html").write_text("\n".join(parts), encoding="utf-8")


def make_presentation(
    output, path, exploration, cv, summary, selected, points, winner_text, comparisons, limits
):
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    notes = []

    def text(slide, content, x, y, width, height, size=22, color="173440", bold=False):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        box.text_frame.word_wrap = True
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = content
        paragraph.font.name = "Liberation Sans"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string(color)
        return box

    def slide(title, subtitle, speaking):
        page = presentation.slides.add_slide(presentation.slide_layouts[6])
        page.background.fill.solid()
        page.background.fill.fore_color.rgb = RGBColor.from_string("F4F3EE")
        text(page, "INFO1185  /  FRAUDETECTION", 0.5, 0.2, 12, 0.3, 11)
        text(page, title, 0.5, 0.72, 12.3, 0.7, 30, bold=True)
        text(page, subtitle, 0.5, 1.5, 12.3, 0.65, 16)
        text(page, f"{len(presentation.slides):02d}", 12.2, 7.05, 0.6, 0.3, 11)
        page.notes_slide.notes_text_frame.text = speaking
        notes.append(f"## {len(presentation.slides)}. {title}\n\n{speaking}\n")
        return page

    def bullets(page, items, size=22):
        for i, item in enumerate(items):
            text(page, item, 0.7, 2.35 + i * 0.95, 11.8, 0.85, size)

    def table(page, frame, y=2.3, height=4.1, size=16, widths=None):
        shape = page.shapes.add_table(
            len(frame) + 1, len(frame.columns), Inches(0.5), Inches(y), Inches(12.3), Inches(height)
        )
        widths = widths or [4.2, *([8.1 / (len(frame.columns) - 1)] * (len(frame.columns) - 1))]
        for column, width in zip(shape.table.columns, widths, strict=True):
            column.width = Inches(width)
        for i, row in enumerate([list(frame.columns), *frame.values.tolist()]):
            for j, value in enumerate(row):
                cell = shape.table.cell(i, j)
                cell.text = str(value)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(
                    "173440" if i == 0 else ("FFFFFF" if i % 2 else "E6ECEC")
                )
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.name = "Liberation Sans"
                    paragraph.font.size = Pt(size)
                    paragraph.font.color.rgb = RGBColor.from_string(
                        "FFFFFF" if i == 0 else "173440"
                    )

    page = slide(
        "Detectar fraude: minimizar lo que se pierde",
        "Comparación de cuatro clasificadores, con y sin compensación del desbalance.",
        "Integrante 1: presentar el objetivo. La sensibilidad de fraude es el criterio "
        "prioritario; las falsas alarmas ayudan a entender el costo. "
        "Duración prevista de las principales: 10–12 minutos.",
    )
    bullets(
        page,
        [
            "k-NN · Árbol de decisión · SVM · Random Forest propio",
            "Prioridad: mayor sensibilidad. Desempate: mayor precisión del fraude.",
            "Ocho experimentos, una partición compartida y evaluación reservada.",
        ],
    )
    page = slide(
        "Pocos fraudes, muchas transacciones",
        f"{exploration['raw_rows']:,} filas originales → "
        f"{exploration['clean_rows']:,} únicas → {exploration['sample_rows']:,} en la muestra.",
        f"Integrante 1: fuente Kaggle, descarga pública TensorFlow. Eliminamos "
        f"{exploration['duplicate_rows_removed']} duplicados exactos para impedir copias "
        f"entre train y test. No hay faltantes. Los V1–V28 ya vienen de PCA; no aplicamos "
        "otro PCA. La muestra estratificada reduce el costo de los modelos.",
    )
    raw_fraction = exploration["raw_class_counts"]["1"] / exploration["raw_rows"]
    sample_fraction = exploration["sample_class_counts"]["1"] / exploration["sample_rows"]
    text(
        page,
        f"Kaggle / copia TensorFlow. 30 predictores: Time, Amount, V1–V28. "
        f"Class: 0 legítima, 1 fraude. Sin faltantes; "
        f"{exploration['duplicate_rows_removed']:,} duplicados eliminados.",
        0.7,
        2.1,
        11.9,
        0.6,
        16,
    )
    text(
        page,
        f"Original: {exploration['raw_class_counts']['1']} fraudes ({100 * raw_fraction:.3f} %). "
        f"Muestra: {exploration['sample_class_counts']['1']} ({100 * sample_fraction:.3f} %), "
        f"legítimas {100 * (1 - sample_fraction):.3f} %. "
        f"Train/test: {exploration['train_class_counts']['1']}/"
        f"{exploration['test_class_counts']['1']} fraudes.",
        0.7,
        2.72,
        11.9,
        0.5,
        16,
    )
    page.shapes.add_picture(
        str(output / "graficos/distribucion.png"),
        Inches(0.7),
        Inches(3.25),
        width=Inches(11.9),
        height=Inches(3.65),
    )
    page = slide(
        "El test se abre después de decidir",
        "Partición estratificada 70/30, semilla 42; cinco pliegues idénticos en entrenamiento.",
        "Integrante 1: explicar que validación y test no se balancean. Cada pliegue ajusta "
        "sus transformaciones solo en su entrenamiento. Las ocho búsquedas terminan y "
        "seleccion.json se guarda antes de evaluar prueba. Ninguna decisión usa test.",
    )
    bullets(
        page,
        [
            "Entrenamiento → cinco pliegues → búsqueda → configuraciones congeladas",
            "Dentro del pliegue: escalado → selección → compensación → clasificador",
            "Reajuste completo → test fijo con tres semillas 42, 43 y 44",
            "CV: desviación entre pliegues. Test: variación entre ajustes, ddof=1.",
        ],
        21,
    )
    page = slide(
        "Seleccionar variables y compensar dentro del pliegue",
        "20 de 30 predictores mediante ANOVA F. Class nunca se incluye en X.",
        "Integrante 1: ANOVA detecta diferencias marginales de medias, pero no todas las "
        "interacciones. k=20 se fijó antes de evaluar. Escalamos k-NN y SVM porque usan "
        "distancias; árboles no lo necesitan. Ratio 0.1 retiene diez legítimas por fraude "
        "y todos los fraudes del entrenamiento del pliegue. Los recuentos están guardados.",
    )
    bullets(
        page,
        [
            "k-NN y SVM se escalan por usar distancias; árboles no dependen de unidades.",
            "ANOVA compara medias marginales; 20 variables acotan el costo. No se aplica otro PCA.",
            "Selección supervisada antes del submuestreo, solo con el train del pliegue.",
            "Con compensación: un fraude por diez transacciones legítimas.",
        ],
        21,
    )
    page = slide(
        "Un bosque construido por la dupla",
        "31 árboles: bootstrap con reemplazo, variables aleatorias y votación mayoritaria.",
        "Integrante 2: recorrer forest.py. Cada árbol recibe n índices con reemplazo, "
        "por eso hay repeticiones. Usamos semillas distintas y max_features=sqrt en cada "
        "división. El bootstrap parte del train disponible, ya compensado si corresponde. "
        "No usamos RandomForestClassifier ni BaggingClassifier. Los votos son etiquetas.",
    )
    bullets(
        page,
        [
            "Por árbol: n filas con reemplazo y semilla propia → un DecisionTreeClassifier.",
            "En cada división: considerar sqrt del número de variables disponibles.",
            "Predicción final: mayoría de etiquetas; 31 árboles evita empates.",
            "Configurables: árboles, profundidad, max_features, hoja mínima y semilla.",
        ],
        21,
    )
    text(
        page,
        "Bagging combina modelos entrenados en bootstrap para reducir variación; "
        "max_features busca reducir la correlación entre árboles.",
        0.7,
        6.25,
        11.9,
        0.6,
        18,
    )
    page = slide(
        "Espacios de búsqueda fijados antes de evaluar",
        "Misma grilla con y sin compensación. Dos candidatos por modelo; 80 ajustes de CV.",
        "Integrante 2: comparar vecindades, profundidad y regularización. Se fijan 31 árboles "
        "y selección de 20 variables para acotar el costo. La grilla es pequeña, por lo que "
        "no se afirma haber encontrado el óptimo global. Ningún valor se eligió mirando test.",
    )
    table(
        page,
        pd.DataFrame(
            [
                ["k-NN", "Vecinos: 3, 7", "Distancia euclídea; comparar vecindad pequeña y mayor."],
                ["Árbol", "Profundidad: 6, libre", "Hoja mínima 2; complejidad acotada y libre."],
                ["SVM", "C: 1, 10", "RBF, gamma=scale; regularización moderada y menor."],
                [
                    "Bosque propio",
                    "Profundidad: 8, libre",
                    "31 árboles, sqrt, hoja mínima 2; diversidad y costo.",
                ],
            ],
            columns=["Modelo", "Candidatos", "Parámetros fijos y motivo"],
        ),
        height=3.5,
        size=18,
        widths=[2.1, 3.3, 6.9],
    )
    text(
        page,
        "Prioridad: sensibilidad media. Empate exacto: precisión media; "
        "si persiste, primer candidato de la grilla.",
        0.7,
        6.25,
        11.9,
        0.6,
        18,
    )
    page = slide(
        "Búsqueda acotada, selección por sensibilidad",
        "Dos candidatos por modelo × cinco pliegues × ocho condiciones = 80 ajustes de CV.",
        "Integrante 2: k-NN compara k=3/7; árbol profundidad 6/libre; SVM RBF C=1/10; "
        "bosque profundidad 8/libre. Árbol y bosque fijan dos ejemplos por hoja. "
        "Estos espacios exploran suavidad, complejidad y regularización a costo acotado. "
        "La tabla usa sensibilidad media y desviación entre pliegues. Empate por precisión.",
    )
    rows = []
    for _, row in cv.iterrows():
        params = json.loads(row.params)
        configuration = ", ".join(f"{k.replace('model__', '')}={v}" for k, v in params.items())
        rows.append([label(row), configuration, percent(row.sensitivity_mean, row.sensitivity_std)])
    table(
        page,
        pd.DataFrame(rows, columns=["Experimento", "Parámetro elegido", "Sensibilidad %"]),
        size=15,
    )
    for title, metrics in [
        (
            "Prueba: detectar fraude y controlar las alertas",
            ("sensitivity", "specificity", "precision"),
        ),
        ("Prueba: F1, exactitud y falsos negativos", ("f1", "accuracy", "fnr")),
    ]:
        page = slide(
            title,
            "Porcentajes: media ± desviación de tres ajustes sobre el mismo test.",
            "Integrante 2: valores calculados con fraude como positivo. Sensibilidad="
            "TP/(TP+FN), especificidad=TN/(TN+FP), precisión=TP/(TP+FP), F1="
            "2TP/(2TP+FP+FN), exactitud=(TP+TN)/N y FNR=1-sensibilidad. "
            "Desviación cero puede indicar un procedimiento determinista; no certeza poblacional.",
        )
        table(page, metric_table(summary, metrics), size=16)
    page = slide(
        "El costo de detectar más",
        "Azul: sin compensación. Naranja: con compensación. Medias de tres ajustes.",
        "Integrante 2: comparar fraudes detectados con falsas alarmas. " + " ".join(points),
    )
    page.shapes.add_picture(
        str(output / "graficos/costos.png"),
        Inches(0.6),
        Inches(2.3),
        width=Inches(12.1),
        height=Inches(4.65),
    )
    page = slide(
        "Árbol frente a bosque; exactitud frente a sensibilidad",
        "Combinar árboles no garantiza mejorar cada métrica en cada condición.",
        "Integrante 2: "
        + " ".join(comparisons)
        + " El bagging busca reducir la variación y max_features reduce la correlación. "
        "La exactitud oculta el desbalance; nunca decide la selección.",
    )
    baseline = load_json(output / "referencias_triviales.json")
    bullets(
        page,
        [
            *comparisons,
            f"Siempre legítimo: exactitud {100 * baseline['always_legitimate']['accuracy']:.3f} %, "
            f"sensibilidad 0 %; {baseline['always_legitimate']['fn']} fraudes omitidos.",
            "Siempre fraude: sensibilidad 100 %; "
            f"{baseline['always_fraud']['fp']:,} falsas alarmas.",
        ],
        19,
    )
    page = slide(
        "Configuración elegida y límites de la evidencia",
        "El ganador se eligió por validación, antes de mirar los resultados de prueba.",
        "Ambos integrantes: "
        + winner_text
        + " "
        + " ".join(limits)
        + " Mejoras futuras: más datos y fraudes, validación temporal y grillas más amplias; "
        "esas mejoras requieren un nuevo protocolo y un test aún no utilizado.",
    )
    text(page, winner_text, 0.7, 2.3, 11.9, 1.15, 21, bold=True)
    concise_limits = [
        f"Solo {exploration['test_class_counts']['1']} fraudes. La desviación entre ajustes "
        "no mide incertidumbre sobre nuevas muestras.",
        "Muestreo y submuestreo pierden información; ANOVA es marginal y la grilla es pequeña.",
        "La división aleatoria no valida cambios temporales ni generalización a otros bancos.",
    ]
    for i, item in enumerate(concise_limits):
        text(page, item, 0.7, 3.65 + i * 0.75, 11.9, 0.65, 18)
    text(
        page,
        "Mejoras futuras: más fraudes, validación temporal y grillas más amplias. "
        "Requieren un nuevo protocolo y un test no utilizado.",
        0.7,
        6.15,
        11.9,
        0.7,
        18,
    )
    # Respaldo: las ocho matrices y las características completas quedan accesibles.
    for group in (list(summary.iloc[:4].iterrows()), list(summary.iloc[4:].iterrows())):
        page = slide(
            "Respaldo: matrices de confusión",
            "Semilla 42. Filas: clase real; columnas: predicción. Fraude es la clase positiva.",
            "Usar si se solicita explicar TP, FN, FP o TN. Las otras semillas y conteos "
            "individuales están en prueba_repeticiones.csv y predicciones.csv.gz.",
        )
        for position, (_, row) in enumerate(group):
            picture = output / "graficos" / f"matriz_{row['model']}_{row['condition']}.png"
            page.shapes.add_picture(
                str(picture),
                Inches(0.9 + (position % 2) * 6.1),
                Inches(2.15 + (position // 2) * 2.5),
                width=Inches(4.25),
                height=Inches(2.45),
            )
    features = load_json(output / "caracteristicas.json")
    feature_rows = [row for row in features if row["seed"] == 42]
    for group in (feature_rows[:4], feature_rows[4:]):
        page = slide(
            "Respaldo: características y recuentos finales",
            "Selección reajustada con todo el entrenamiento, semilla 42.",
            "ANOVA se ajusta antes de remuestrear. Con el mismo train y k=20, la selección "
            "puede ser idéntica para todos los modelos. Las tres semillas están registradas.",
        )
        for i, row in enumerate(group):
            content = (
                f"{label(row)}: "
                + ", ".join(row["selected_features"])
                + f"\nTrain antes {row['before']}; después {row['after']}."
            )
            text(page, content, 0.7, 2.25 + i * 1.1, 11.9, 1.0, 16)
    cover = presentation.slides[0]
    text(cover, "\n".join(AUTHORS), 0.7, 5.9, 11.9, 0.9, 22, bold=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(path)
    (path.parent / "guion_defensa.md").write_text(
        "# Guion de defensa\n\n12 diapositivas principales y cuatro de respaldo. "
        "Fernando Alonso Valdés Gómez y Juan Francisco Muñoz Veloso. "
        "Distribuir 5–6 minutos por persona.\n\n" + "\n".join(notes),
        encoding="utf-8",
    )


def generate_report(output, presentation_path):
    output, presentation_path = Path(output), Path(presentation_path)
    protocol = load_json(output / "protocolo.json")
    if protocol["status"] != "complete":
        raise ValueError("Solo se generan documentos de una ejecución completa.")
    exploration = load_json(output / "exploracion.json")
    selected = load_json(output / "seleccion.json")
    cv = pd.read_csv(output / "validacion.csv")
    summary = pd.read_csv(output / "prueba_resumen.csv")
    repetitions = pd.read_csv(output / "prueba_repeticiones.csv")
    proportions = []
    for stage in ("raw", "clean", "sample", "train", "test"):
        counts = exploration[f"{stage}_class_counts"]
        total = sum(counts.values())
        proportions.append(
            {
                "stage": stage,
                "rows": total,
                "legitimate_count": counts["0"],
                "fraud_count": counts["1"],
                "legitimate_fraction": counts["0"] / total,
                "fraud_fraction": counts["1"] / total,
            }
        )
    pd.DataFrame(proportions).to_csv(output / "proporciones_clases.csv", index=False)
    points, winner_text, comparisons, limits = analysis_points(summary, selected, cv, exploration)
    make_figures(output, exploration, cv, summary, repetitions)
    make_html(output, exploration, cv, summary, selected, points, winner_text, comparisons, limits)
    make_presentation(
        output,
        presentation_path,
        exploration,
        cv,
        summary,
        selected,
        points,
        winner_text,
        comparisons,
        limits,
    )
    (output / "analisis.md").write_text(
        "# Análisis de resultados\n\n"
        + winner_text
        + "\n\n"
        + "\n\n".join([*points, *comparisons, *limits]),
        encoding="utf-8",
    )
